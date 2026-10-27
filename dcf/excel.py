"""Write the valuation as an Excel workbook built from live formulas.

Convention (standard in banking models): blue font = hard-coded input pulled from a source,
black font = formula. Every number downstream of the inputs is a formula, so changing an input
(say terminal growth on the Inputs sheet) re-prices the whole model in Excel.
Dollar amounts are in $ millions; shares in millions; per-share values in $.
"""
from __future__ import annotations

from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config
from .pipeline import Model

BLUE = Font(color="0000CC")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
HEAD_FILL = PatternFill("solid", fgColor="DDE4EE")
KEY_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Border(top=Side(style="thin"))

MIL = '#,##0;(#,##0);"-"'
MIL1 = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%)'
PCT2 = '0.00%;(0.00%)'
USD = '$#,##0.00'
MULT = '0.0"x"'
DATE = 'yyyy-mm-dd'
NUM2 = '0.00'
NUM3 = '0.000'

M = 1e6


def _a(sheet: str, col: int, row: int, absolute: bool = True) -> str:
    c = get_column_letter(col)
    return f"'{sheet}'!${c}${row}" if absolute else f"'{sheet}'!{c}{row}"


class Rows:
    """Append label/value rows to a sheet and remember where each key landed."""

    def __init__(self, ws, start_row: int = 3, value_col: int = 2):
        self.ws, self.row, self.col, self.at = ws, start_row, value_col, {}

    def add(self, key, label, value, fmt=None, note="", is_input=False, bold=False, fill=None):
        ws, r = self.ws, self.row
        ws.cell(r, 1, label).font = BOLD if bold else Font()
        c = ws.cell(r, self.col, value)
        if fmt:
            c.number_format = fmt
        if is_input:
            c.font = BLUE
        elif bold:
            c.font = BOLD
        if fill:
            c.fill = fill
            ws.cell(r, 1).fill = fill
        if note:
            ws.cell(r, self.col + 1, note).alignment = Alignment(wrap_text=False)
        if key:
            self.at[key] = _a(ws.title, self.col, r)
        self.row += 1
        return r

    def gap(self, n=1):
        self.row += n

    def header(self, text):
        c = self.ws.cell(self.row, 1, text)
        c.font = BOLD
        c.fill = HEAD_FILL
        for j in range(2, self.col + 2):
            self.ws.cell(self.row, j).fill = HEAD_FILL
        self.row += 1


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_workbook(m: Model, path) -> dict:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    ref: dict[str, str] = {}

    # ------------------------------------------------------------------ Inputs
    ws = wb.create_sheet("Inputs")
    ws["A1"] = f"{m.company} ({m.ticker}) - source data"
    ws["A1"].font = TITLE
    ws["A2"] = "Blue = hard-coded input from the cited source. $ millions unless stated."
    R = Rows(ws)
    R.header("Company and market data")
    R.add("company", "Company", m.company, note=f"SEC CIK {m.cik}", is_input=True)
    R.add("ticker", "Ticker", m.ticker, is_input=True)
    R.add("valdate", "Valuation date (date of share price)", date.fromisoformat(m.quote["date"]),
          DATE, m.quote["source"], True)
    R.add("price", "Share price ($)", m.quote["price"], USD, m.quote["source"], True)
    R.add("shares", "Shares outstanding (m)", m.shares / M, MIL1,
          f"10-K/10-Q cover page, dei:EntityCommonStockSharesOutstanding as of {m.shares_date}", True)
    R.add("week52_low", "52-week low close ($)", m.week52[0], USD, "Yahoo Finance via yfinance", True)
    R.add("week52_high", "52-week high close ($)", m.week52[1], USD, "Yahoo Finance via yfinance", True)
    R.gap()
    R.header(f"Latest balance sheet ({m.bs.as_of}, most recent 10-K or 10-Q)")
    comp = m.bs.components
    R.add("bs_date", "Balance sheet date", m.bs.as_of, DATE, "SEC companyfacts", True)
    R.add("cash", "Cash and equivalents", m.bs.cash / M, MIL, comp["cash"]["tag"] or "", True)
    R.add("sti", "Short-term investments", m.bs.st_investments / M, MIL,
          comp["st_investments"]["tag"] or "not reported (0)", True)
    R.add("debt_cur", "Debt - current portion", (comp["debt_current"]["value"]) / M, MIL,
          comp["debt_current"]["tag"] or "not reported (0)", True)
    R.add("debt_st", "Short-term borrowings", (comp["st_borrowings"]["value"]) / M, MIL,
          comp["st_borrowings"]["tag"] or "not reported (0)", True)
    R.add("debt_lt", "Debt - non-current", (comp["debt_noncurrent"]["value"]) / M, MIL,
          comp["debt_noncurrent"]["tag"] or "not reported (0)", True)
    R.add("debt", "Total debt", f"={R.at['debt_cur']}+{R.at['debt_st']}+{R.at['debt_lt']}", MIL,
          "Operating lease liabilities excluded: lease cost is already inside EBIT", bold=True)
    R.add("nci", "Noncontrolling interest", m.bs.noncontrolling_interest / M, MIL,
          comp["noncontrolling_interest"]["tag"] or "not reported (0)", True)
    R.gap()
    R.header("Rates")
    R.add("rf", "Risk-free rate (10-year Treasury)", m.rf["rate"], PCT2,
          f"{m.rf['source']}, {m.rf['date']}", True)
    R.add("erp", "Equity risk premium", m.erp["erp"], PCT2,
          f"{m.erp['source']}, {m.erp['date']}", True)
    R.add("tg", "Terminal growth rate (choice)", m.assumptions.terminal_growth, PCT2,
          "Long-run nominal growth: ~2% inflation plus modest real growth; below the risk-free rate",
          True, fill=KEY_FILL)
    R.add("lastfye", "Last reported fiscal year end", m.history.fiscal_year_ends[-1], DATE,
          "Latest 10-K", True)
    ref.update(R.at)
    R.gap()
    R.header(f"Synthetic rating table (Damodaran, large non-financial firms, {m.ratings[0]['as_of']})")
    hdr = R.row
    for j, h in enumerate(["Coverage from", "Coverage to", "Rating", "Default spread"], start=1):
        ws.cell(hdr, j, h).font = BOLD
    first = hdr + 1
    for i, row in enumerate(m.ratings):
        r = first + i
        for j, (v, f) in enumerate([(row["low"], NUM2), (row["high"], NUM2),
                                    (row["rating"], None), (round(row["spread"], 6), PCT2)],
                                   start=1):
            c = ws.cell(r, j, v)
            c.font = BLUE
            if f:
                c.number_format = f
    last = first + len(m.ratings) - 1
    ref["rt_low"] = f"'Inputs'!$A${first}:$A${last}"
    ref["rt_rating"] = f"'Inputs'!$C${first}:$C${last}"
    ref["rt_spread"] = f"'Inputs'!$D${first}:$D${last}"
    _widths(ws, [44, 16, 90, 14])

    # ------------------------------------------------------------------ Historicals
    hs = wb.create_sheet("Historicals")
    t = m.history.table
    years = list(t.columns)
    ny = len(years)
    hs["A1"] = f"{m.company} - fiscal-year history from 10-K XBRL ($ millions)"
    hs["A1"].font = TITLE
    hs["A2"] = ("Fiscal year = calendar year in which the year ends. Blue = as reported in XBRL; "
                "black = formula.")
    hs.cell(3, 1, "Line item").font = BOLD
    for j, y in enumerate(years):
        hs.cell(3, 2 + j, f"FY{y}").font = BOLD
        hs.cell(3, 2 + j).fill = HEAD_FILL
    avg_col = 2 + ny
    src_col = avg_col + 1
    hs.cell(3, avg_col, f"Avg last {config.AVERAGING_YEARS}").font = BOLD
    hs.cell(3, src_col, "XBRL tag / definition").font = BOLD
    hrow: dict[str, int] = {}
    r = 4

    def col(j):
        return get_column_letter(2 + j)

    raw_items = [
        ("revenue", "Revenue", MIL, M),
        ("operating_income", "Operating income (reported; 0 = not tagged)", MIL, M),
        ("pretax_income", "Pre-tax income", MIL, M),
        ("interest_expense", "Interest expense", MIL, M),
        ("interest_income", "Interest income", MIL, M),
        ("da", "Depreciation & amortization", MIL, M),
        ("income_tax", "Income tax expense", MIL, M),
        ("capex", "Capital expenditures", MIL, M),
        ("net_income", "Net income", MIL, M),
        ("eps_diluted", "Diluted EPS ($)", USD, 1),
        ("diluted_shares", "Diluted weighted shares (m)", MIL1, M),
        ("current_assets", "Total current assets", MIL, M),
        ("cash", "Cash and equivalents", MIL, M),
        ("st_investments", "Short-term investments", MIL, M),
        ("current_liabilities", "Total current liabilities", MIL, M),
        ("debt_current", "Current debt", MIL, M),
        ("st_borrowings", "Short-term borrowings", MIL, M),
        ("lease_liab_current", "Current operating lease liabilities", MIL, M),
        ("lease_liab_total", "Total operating lease liabilities (memo, not in net debt)", MIL, M),
    ]
    for key, label, fmt, scale in raw_items:
        hs.cell(r, 1, label)
        for j, y in enumerate(years):
            v = t.at[key, y] if key in t.index else None
            v = 0.0 if (v is None or pd.isna(v)) else float(v) / scale
            c = hs.cell(r, 2 + j, v)
            c.number_format, c.font = fmt, BLUE
        hs.cell(r, src_col, m.history.sources.get(key, "not reported"))
        hrow[key] = r
        r += 1
    r += 1
    hs.cell(r, 1, "Derived").font = BOLD
    r += 1

    def frow(key, label, fn, fmt, avg=False, note=""):
        nonlocal r
        hs.cell(r, 1, label)
        for j in range(ny):
            f = fn(j)
            if f is None:
                continue
            c = hs.cell(r, 2 + j, f)
            c.number_format = fmt
        if avg:
            a = f"=AVERAGE({col(ny - config.AVERAGING_YEARS)}{r}:{col(ny - 1)}{r})"
            c = hs.cell(r, avg_col, a)
            c.number_format, c.font = fmt, BOLD
        hs.cell(r, src_col, note)
        hrow[key] = r
        r += 1

    h = hrow
    frow("ebit", "EBIT", lambda j: f"=IF({col(j)}{h['operating_income']}<>0,"
         f"{col(j)}{h['operating_income']},{col(j)}{h['pretax_income']}+"
         f"{col(j)}{h['interest_expense']}-{col(j)}{h['interest_income']})", MIL,
         note="Reported operating income; if not tagged: pre-tax + interest expense - interest income")
    frow("ebitda", "EBITDA", lambda j: f"={col(j)}{h['ebit']}+{col(j)}{h['da']}", MIL)
    frow("growth", "Revenue growth", lambda j: None if j == 0 else
         f"={col(j)}{h['revenue']}/{col(j - 1)}{h['revenue']}-1", PCT)
    frow("ebit_margin", "EBIT margin", lambda j: f"={col(j)}{h['ebit']}/{col(j)}{h['revenue']}",
         PCT, True)
    frow("tax_rate", "Effective tax rate", lambda j:
         f"={col(j)}{h['income_tax']}/{col(j)}{h['pretax_income']}", PCT, True)
    frow("da_pct", "D&A % of revenue", lambda j: f"={col(j)}{h['da']}/{col(j)}{h['revenue']}",
         PCT2, True)
    frow("capex_pct", "Capex % of revenue", lambda j:
         f"={col(j)}{h['capex']}/{col(j)}{h['revenue']}", PCT2, True)
    frow("nwc", "Net working capital (operating)", lambda j:
         f"=({col(j)}{h['current_assets']}-{col(j)}{h['cash']}-{col(j)}{h['st_investments']})-"
         f"({col(j)}{h['current_liabilities']}-{col(j)}{h['debt_current']}-"
         f"{col(j)}{h['st_borrowings']}-{col(j)}{h['lease_liab_current']})", MIL,
         note="(current assets - cash - ST investments) - (current liabilities - current debt - "
              "ST borrowings - current lease liabilities)")
    frow("nwc_pct", "NWC % of revenue", lambda j: f"={col(j)}{h['nwc']}/{col(j)}{h['revenue']}",
         PCT2, True)
    frow("ufcf", "Unlevered free cash flow (historical, for reference)", lambda j: None if j == 0
         else f"={col(j)}{h['ebit']}*(1-{col(j)}{h['tax_rate']})+{col(j)}{h['da']}-"
              f"{col(j)}{h['capex']}-({col(j)}{h['nwc']}-{col(j - 1)}{h['nwc']})", MIL,
         note="EBIT x (1 - tax) + D&A - capex - change in NWC")
    lastc = col(ny - 1)
    hist_last = lambda key: f"'Historicals'!${lastc}${h[key]}"
    hist_avg = lambda key: f"'Historicals'!${get_column_letter(avg_col)}${h[key]}"
    hs.freeze_panes = "B4"
    _widths(hs, [52] + [12] * ny + [13, 70])
    for note_i, n in enumerate(m.history.notes):
        hs.cell(r + 1 + note_i, 1, "Note: " + n)

    # ------------------------------------------------------------------ Beta
    bsh = wb.create_sheet("Beta")
    bsh["A1"] = f"Beta: {m.ticker} vs S&P 500, monthly returns"
    bsh["A1"].font = TITLE
    bsh["A2"] = ("Adjusted month-end closes (Yahoo Finance via yfinance; index = ^GSPC). "
                 "Beta = SLOPE of stock returns on index returns.")
    for j, hdr_ in enumerate(["Month", f"{m.ticker} adj. close", "S&P 500", f"{m.ticker} return",
                              "S&P 500 return"], start=1):
        bsh.cell(4, j, hdr_).font = BOLD
    bt = m.beta_table
    for i, (idx, row) in enumerate(bt.iterrows()):
        rr = 5 + i
        bsh.cell(rr, 1, idx.date()).number_format = DATE
        c = bsh.cell(rr, 2, float(row["stock"])); c.font = BLUE; c.number_format = NUM2
        c = bsh.cell(rr, 3, float(row["market"])); c.font = BLUE; c.number_format = NUM2
        if i > 0:
            bsh.cell(rr, 4, f"=B{rr}/B{rr - 1}-1").number_format = PCT2
            bsh.cell(rr, 5, f"=C{rr}/C{rr - 1}-1").number_format = PCT2
    first_ret, last_ret = 6, 5 + len(bt) - 1
    bsh["G4"], bsh["G5"], bsh["G6"], bsh["G7"] = "Beta (SLOPE)", "R-squared", "Observations", \
        "Blume-adjusted beta (0.67 x raw + 0.33), for reference only"
    bsh["H4"] = f"=SLOPE(D{first_ret}:D{last_ret},E{first_ret}:E{last_ret})"
    bsh["H5"] = f"=RSQ(D{first_ret}:D{last_ret},E{first_ret}:E{last_ret})"
    bsh["H6"] = f"=COUNT(D{first_ret}:D{last_ret})"
    bsh["H7"] = "=0.67*H4+0.33"
    for c_ in ("H4", "H5", "H7"):
        bsh[c_].number_format = NUM3
    bsh["H4"].font = BOLD
    ref["beta"] = "'Beta'!$H$4"
    _widths(bsh, [12, 16, 12, 14, 14, 3, 52, 10])

    # ------------------------------------------------------------------ Comps
    cs = wb.create_sheet("Comps")
    cs["A1"] = "Trading comparables (last fiscal year; prices as of the valuation date)"
    cs["A1"].font = TITLE
    cs["A2"] = ("EV = market cap + debt - cash - ST investments + NCI (latest 10-K/10-Q). EBITDA = "
                "EBIT + D&A and EPS = diluted EPS, both from the latest 10-K. $ millions.")
    heads = ["Ticker", "Company", "FY", "Price ($)", "Shares (m)", "Market cap", "Debt",
             "Cash + STI", "NCI", "Enterprise value", "EBITDA", "Diluted EPS ($)", "EV / EBITDA",
             "P / E"]
    for j, hh in enumerate(heads, start=1):
        cs.cell(4, j, hh).font = BOLD
        cs.cell(4, j).fill = HEAD_FILL
    prow0 = 5

    def comp_row(rr, p, linked=False):
        vals = [p.ticker, p.name, f"FY{p.fiscal_year}"]
        for j, v in enumerate(vals, start=1):
            cs.cell(rr, j, v)
        if linked:
            cells = [f"={ref['price']}", f"={ref['shares']}", None, f"={ref['debt']}",
                     f"={ref['cash']}+{ref['sti']}", f"={ref['nci']}", None,
                     f"={hist_last('ebitda')}", f"={hist_last('eps_diluted')}"]
        else:
            cells = [p.price, p.shares / M, None, p.debt / M, p.cash / M,
                     p.noncontrolling_interest / M, None, p.ebitda / M, p.eps]
        fmts = [USD, MIL1, MIL, MIL, MIL, MIL, MIL, MIL, USD]
        for j, (v, f) in enumerate(zip(cells, fmts), start=4):
            if v is None:
                continue
            c = cs.cell(rr, j, v)
            c.number_format = f
            if not linked:
                c.font = BLUE
        cs.cell(rr, 6, f"=D{rr}*E{rr}").number_format = MIL
        cs.cell(rr, 10, f"=F{rr}+G{rr}-H{rr}+I{rr}").number_format = MIL
        cs.cell(rr, 13, f"=IF(K{rr}>0,J{rr}/K{rr},NA())").number_format = MULT
        cs.cell(rr, 14, f"=IF(L{rr}>0,D{rr}/L{rr},NA())").number_format = MULT

    for i, p in enumerate(m.peers):
        comp_row(prow0 + i, p)
    plast = prow0 + len(m.peers) - 1
    stat = plast + 1
    for k, fn in enumerate(["MEDIAN", "AVERAGE", "MIN", "MAX"]):
        rr = stat + k
        cs.cell(rr, 2, f"Peer {fn.lower() if fn != 'AVERAGE' else 'mean'}").font = BOLD
        for cc in ("M", "N"):
            c = cs[f"{cc}{rr}"]
            c.value = f"={fn}({cc}{prow0}:{cc}{plast})"
            c.number_format = MULT
            c.font = BOLD
    med_row, min_row, max_row = stat, stat + 2, stat + 3
    trow = stat + 5
    cs.cell(trow - 1, 1, "Target").font = BOLD
    comp_row(trow, m.target_comp, linked=True)
    ref["exit_multiple"] = f"'Comps'!$M${med_row}"

    ir = trow + 3
    cs.cell(ir, 1, "Implied value per share of target").font = BOLD
    cs.cell(ir, 1).fill = HEAD_FILL
    for j, hh in enumerate(["", "Low (peer min)", "Median", "High (peer max)"], start=2):
        cs.cell(ir, j + 2, hh).font = BOLD
    cs.cell(ir + 1, 1, "From EV / EBITDA")
    cs.cell(ir + 2, 1, "From P / E")
    for j, srow in enumerate([min_row, med_row, max_row]):
        cc = get_column_letter(5 + j)
        cs[f"{cc}{ir + 1}"] = (f"=(M{srow}*K{trow}-G{trow}+H{trow}-I{trow})/E{trow}")
        cs[f"{cc}{ir + 2}"] = f"=N{srow}*L{trow}"
        cs[f"{cc}{ir + 1}"].number_format = USD
        cs[f"{cc}{ir + 2}"].number_format = USD
    ref["comps_ev_med"] = f"'Comps'!$F${ir + 1}"
    ref["comps_pe_med"] = f"'Comps'!$F${ir + 2}"
    _widths(cs, [9, 38, 9, 11, 11, 12, 10, 11, 8, 15, 11, 14, 12, 9])

    # ------------------------------------------------------------------ WACC
    wsw = wb.create_sheet("WACC")
    wsw["A1"] = "Weighted average cost of capital"
    wsw["A1"].font = TITLE
    W = Rows(wsw)
    W.header("Cost of equity (CAPM)")
    W.add("rf", "Risk-free rate", f"={ref['rf']}", PCT2, "Inputs")
    W.add("beta", "Levered beta (5y monthly regression)", f"={ref['beta']}", NUM3, "Beta sheet")
    W.add("erp", "Equity risk premium", f"={ref['erp']}", PCT2, "Inputs (Damodaran implied ERP)")
    W.add("ke", "Cost of equity = Rf + Beta x ERP", f"={W.at['rf']}+{W.at['beta']}*{W.at['erp']}",
          PCT2, bold=True)
    W.gap()
    W.header("Cost of debt (synthetic rating from interest coverage)")
    W.add("ebit", "EBIT, last fiscal year", f"={hist_last('ebit')}", MIL, "Historicals")
    W.add("int", "Interest expense, last fiscal year", f"={hist_last('interest_expense')}", MIL,
          "Historicals")
    W.add("cov", "Interest coverage (EBIT / interest)",
          f"=IF({W.at['int']}>0,{W.at['ebit']}/{W.at['int']},100000)", '0.0"x"')
    W.add("rating", "Synthetic rating", f"=LOOKUP({W.at['cov']},{ref['rt_low']},{ref['rt_rating']})",
          None, "Damodaran coverage table (Inputs)")
    W.add("spread", "Default spread", f"=LOOKUP({W.at['cov']},{ref['rt_low']},{ref['rt_spread']})",
          PCT2)
    W.add("kd", "Pre-tax cost of debt = Rf + spread", f"={W.at['rf']}+{W.at['spread']}", PCT2)
    W.add("tax", "Tax rate (3-yr average effective)", f"={hist_avg('tax_rate')}", PCT)
    W.add("kd_at", "After-tax cost of debt", f"={W.at['kd']}*(1-{W.at['tax']})", PCT2, bold=True)
    W.gap()
    W.header("Capital structure (market value of equity, book value of debt)")
    W.add("e", "Equity market value", f"={ref['price']}*{ref['shares']}", MIL)
    W.add("d", "Debt", f"={ref['debt']}", MIL)
    W.add("we", "Equity weight", f"={W.at['e']}/({W.at['e']}+{W.at['d']})", PCT)
    W.add("wd", "Debt weight", f"={W.at['d']}/({W.at['e']}+{W.at['d']})", PCT)
    W.gap()
    W.add("wacc", "WACC", f"={W.at['we']}*{W.at['ke']}+{W.at['wd']}*{W.at['kd_at']}", PCT2,
          bold=True, fill=KEY_FILL)
    ref["wacc"] = W.at["wacc"]
    ref["ke"] = W.at["ke"]
    _widths(wsw, [46, 14, 50])

    # ------------------------------------------------------------------ Assumptions
    asn = wb.create_sheet("Assumptions")
    asn["A1"] = "Forecast assumptions (each is a formula pointing at its source)"
    asn["A1"].font = TITLE
    A = Rows(asn)
    n = config.AVERAGING_YEARS
    first_y = col(ny - 1 - n)
    lo, hi = config.YEAR1_GROWTH_BOUNDS
    A.header("Revenue")
    A.add("cagr", f"Revenue CAGR, FY{years[-1 - n]}-FY{years[-1]}",
          f"=('Historicals'!${lastc}${h['revenue']}/'Historicals'!${first_y}${h['revenue']})^(1/{n})-1",
          PCT2, "Historicals")
    A.add("g1", "Year-1 revenue growth",
          f"=MIN(MAX({A.at['cagr']},{lo}),{hi})", PCT2,
          f"Historical CAGR, bounded to [{lo:.0%}, {hi:.0%}]", fill=KEY_FILL)
    A.add("tg", "Terminal growth", f"={ref['tg']}", PCT2,
          "Inputs. Years 2-5 fade in a straight line from year-1 growth toward this rate")
    A.add("years", "Projection years", config.PROJECTION_YEARS, "0", "", True)
    A.header(f"Margins and reinvestment ({n}-year averages, FY{years[-n]}-FY{years[-1]})")
    A.add("margin", "EBIT margin", f"={hist_avg('ebit_margin')}", PCT2, "Historicals avg",
          fill=KEY_FILL)
    A.add("tax", "Tax rate", f"={hist_avg('tax_rate')}", PCT2, "Historicals avg (effective rate)")
    A.add("da", "D&A % of revenue", f"={hist_avg('da_pct')}", PCT2, "Historicals avg")
    A.add("capex", "Capex % of revenue", f"={hist_avg('capex_pct')}", PCT2, "Historicals avg")
    A.add("nwc", "NWC % of revenue", f"={hist_avg('nwc_pct')}", PCT2,
          "Historicals avg; change in NWC = NWC% x change in revenue")
    A.header("Discounting and terminal value")
    A.add("wacc", "WACC", f"={ref['wacc']}", PCT2, "WACC sheet", fill=KEY_FILL)
    A.add("exit", "Exit EV/EBITDA multiple", f"={ref['exit_multiple']}", MULT,
          "Peer median EV/EBITDA (Comps sheet)")
    A.add(None, "Cash-flow timing", "mid-year", None,
          "Year 1 counts only the part of the fiscal year after the valuation date; terminal "
          "value discounted from the end of year 5")
    ref.update({f"as_{k}": v for k, v in A.at.items()})
    _widths(asn, [44, 14, 90])

    # ------------------------------------------------------------------ DCF
    d = wb.create_sheet("DCF")
    d["A1"] = f"{m.company} - unlevered DCF ($ millions)"
    d["A1"].font = TITLE
    P = config.PROJECTION_YEARS
    pc = [get_column_letter(3 + i) for i in range(P)]  # projection columns C..G
    lc = pc[-1]
    d.cell(3, 2, f"FY{years[-1]}A").font = BOLD
    for i, c_ in enumerate(pc):
        d[f"{c_}3"] = f"FY{int(years[-1]) + i + 1}E"
        d[f"{c_}3"].font = BOLD
    for c_ in ["A", "B"] + pc:
        d[f"{c_}3"].fill = HEAD_FILL
    rows = {}
    labels = [("yr", "Year #"), ("fye", "Fiscal year end"), ("g", "Revenue growth"),
              ("rev", "Revenue"), ("mg", "EBIT margin"), ("ebit", "EBIT"), ("tax", "Tax rate"),
              ("nopat", "NOPAT = EBIT x (1 - tax)"), ("dap", "D&A % of revenue"),
              ("da", "Plus: D&A"), ("cxp", "Capex % of revenue"), ("cx", "Less: capex"),
              ("nwp", "NWC % of revenue"), ("dnwc", "Less: increase in NWC"),
              ("fcf", "Unlevered free cash flow"), ("ebitda", "EBITDA"), (None, ""),
              ("frac", "Share of year after valuation date"),
              ("per", "Discount period (mid-year, years)"), ("df", "Discount factor"),
              ("pv", "PV of free cash flow")]
    for i, (k, lab) in enumerate(labels):
        rr = 4 + i
        d.cell(rr, 1, lab)
        if k:
            rows[k] = rr
    R_ = rows
    d[f"B{R_['fye']}"] = f"={ref['lastfye']}"
    d[f"B{R_['fye']}"].number_format = DATE
    d[f"B{R_['rev']}"] = f"={hist_last('revenue')}"
    d[f"B{R_['rev']}"].number_format = MIL
    for i, c_ in enumerate(pc):
        prev = "B" if i == 0 else pc[i - 1]
        f = {
            "yr": i + 1,
            "fye": f"=EDATE($B${R_['fye']},12*{c_}{R_['yr']})",
            "g": f"={ref['as_g1']}+({ref['as_tg']}-{ref['as_g1']})*({c_}{R_['yr']}-1)/{ref['as_years']}",
            "rev": f"={prev}{R_['rev']}*(1+{c_}{R_['g']})",
            "mg": f"={ref['as_margin']}",
            "ebit": f"={c_}{R_['rev']}*{c_}{R_['mg']}",
            "tax": f"={ref['as_tax']}",
            "nopat": f"={c_}{R_['ebit']}*(1-{c_}{R_['tax']})",
            "dap": f"={ref['as_da']}",
            "da": f"={c_}{R_['rev']}*{c_}{R_['dap']}",
            "cxp": f"={ref['as_capex']}",
            "cx": f"={c_}{R_['rev']}*{c_}{R_['cxp']}",
            "nwp": f"={ref['as_nwc']}",
            "dnwc": f"={c_}{R_['nwp']}*({c_}{R_['rev']}-{prev}{R_['rev']})",
            "fcf": f"={c_}{R_['nopat']}+{c_}{R_['da']}-{c_}{R_['cx']}-{c_}{R_['dnwc']}",
            "ebitda": f"={c_}{R_['ebit']}+{c_}{R_['da']}",
            "frac": (f"=MAX(0,MIN(1,({c_}{R_['fye']}-{ref['valdate']})/365))" if i == 0 else 1),
            "per": (f"={c_}{R_['frac']}/2" if i == 0 else
                    f"=$C${R_['frac']}+({c_}{R_['yr']}-1)-0.5"),
            "df": f"=1/(1+{ref['as_wacc']})^{c_}{R_['per']}",
            "pv": f"={c_}{R_['fcf']}*{c_}{R_['frac']}*{c_}{R_['df']}",
        }
        fmts = {"yr": "0", "fye": DATE, "g": PCT2, "mg": PCT2, "tax": PCT2, "dap": PCT2,
                "cxp": PCT2, "nwp": PCT2, "frac": NUM3, "per": NUM3, "df": "0.0000"}
        for k, v in f.items():
            c = d[f"{c_}{R_[k]}"]
            c.value = v
            c.number_format = fmts.get(k, MIL)
    for k in ("fcf", "pv"):
        for c_ in ["A"] + pc:
            d[f"{c_}{R_[k]}"].font = BOLD
    for c_ in ["A"] + pc:
        d[f"{c_}{R_['fcf']}"].border = THIN

    v0 = R_["pv"] + 2
    d.cell(v0, 1, "Valuation").font = BOLD
    d.cell(v0, 2, "Gordon growth").font = BOLD
    d.cell(v0, 3, "Exit multiple").font = BOLD
    for c_ in "ABC":
        d[f"{c_}{v0}"].fill = HEAD_FILL
    vr = {}
    vlabels = [("sum", "Sum of PV of free cash flows"), ("tv", "Terminal value (end of year 5)"),
               ("tvp", "Terminal value discount period (years)"), ("pvtv", "PV of terminal value"),
               ("ev", "Enterprise value"), ("debt", "Less: debt"), ("cash", "Plus: cash and ST investments"),
               ("nci", "Less: noncontrolling interest"), ("eq", "Equity value"),
               ("sh", "Shares outstanding (m)"), ("px", "Implied value per share ($)"),
               ("mkt", "Current share price ($)"), ("up", "Upside / (downside)"),
               ("tvpct", "PV of terminal value as % of EV"),
               ("cross", "Cross-check: implied exit multiple / implied perpetual growth")]
    for i, (k, lab) in enumerate(vlabels):
        rr = v0 + 1 + i
        d.cell(rr, 1, lab)
        vr[k] = rr
    wacc_ = ref["as_wacc"]
    tg_ = ref["as_tg"]
    d[f"B{vr['tv']}"] = f"={lc}{R_['fcf']}*(1+{tg_})/({wacc_}-{tg_})"
    d[f"C{vr['tv']}"] = f"={lc}{R_['ebitda']}*{ref['as_exit']}"
    for c_ in "BC":
        d[f"{c_}{vr['sum']}"] = f"=SUM({pc[0]}{R_['pv']}:{lc}{R_['pv']})"
        d[f"{c_}{vr['tvp']}"] = f"=$C${R_['frac']}+({lc}{R_['yr']}-1)"
        d[f"{c_}{vr['pvtv']}"] = f"={c_}{vr['tv']}/(1+{wacc_})^{c_}{vr['tvp']}"
        d[f"{c_}{vr['ev']}"] = f"={c_}{vr['sum']}+{c_}{vr['pvtv']}"
        d[f"{c_}{vr['debt']}"] = f"=-{ref['debt']}"
        d[f"{c_}{vr['cash']}"] = f"={ref['cash']}+{ref['sti']}"
        d[f"{c_}{vr['nci']}"] = f"=-{ref['nci']}"
        d[f"{c_}{vr['eq']}"] = f"=SUM({c_}{vr['ev']}:{c_}{vr['nci']})"
        d[f"{c_}{vr['sh']}"] = f"={ref['shares']}"
        d[f"{c_}{vr['px']}"] = f"={c_}{vr['eq']}/{c_}{vr['sh']}"
        d[f"{c_}{vr['mkt']}"] = f"={ref['price']}"
        d[f"{c_}{vr['up']}"] = f"={c_}{vr['px']}/{c_}{vr['mkt']}-1"
        d[f"{c_}{vr['tvpct']}"] = f"={c_}{vr['pvtv']}/{c_}{vr['ev']}"
        for k, fmt in [("sum", MIL), ("tv", MIL), ("tvp", NUM3), ("pvtv", MIL), ("ev", MIL),
                       ("debt", MIL), ("cash", MIL), ("nci", MIL), ("eq", MIL), ("sh", MIL1),
                       ("px", USD), ("mkt", USD), ("up", PCT), ("tvpct", PCT)]:
            d[f"{c_}{vr[k]}"].number_format = fmt
        d[f"{c_}{vr['px']}"].font = BOLD
        d[f"{c_}{vr['px']}"].fill = KEY_FILL
    d[f"B{vr['cross']}"] = f"=B{vr['tv']}/{lc}{R_['ebitda']}"
    d[f"B{vr['cross']}"].number_format = MULT
    d[f"C{vr['cross']}"] = (f"=(C{vr['tv']}*{wacc_}-{lc}{R_['fcf']})/(C{vr['tv']}+{lc}{R_['fcf']})")
    d[f"C{vr['cross']}"].number_format = PCT2
    d.freeze_panes = "B4"
    _widths(d, [58, 15, 15, 13, 13, 13, 13])
    ref["px_gordon"] = f"'DCF'!$B${vr['px']}"
    ref["px_exit"] = f"'DCF'!$C${vr['px']}"
    ref["ev_gordon"] = f"'DCF'!$B${vr['ev']}"
    ref["tvpct"] = f"'DCF'!$B${vr['tvpct']}"

    # ------------------------------------------------------------------ Sensitivity
    s = wb.create_sheet("Sensitivity")
    s["A1"] = "Value per share sensitivity (live formulas; projected cash flows held fixed)"
    s["A1"].font = TITLE

    def pv_terms(wcell):
        return "+".join(f"'DCF'!${c_}${R_['fcf']}*'DCF'!${c_}${R_['frac']}/(1+{wcell})^'DCF'!${c_}${R_['per']}"
                        for c_ in pc)

    bridge = f"-{ref['debt']}+{ref['cash']}+{ref['sti']}-{ref['nci']}"
    tvper = f"'DCF'!$B${vr['tvp']}"

    def grid(top, title, col_label, col_vals_formula, cell_formula, col_fmt):
        s.cell(top, 1, title).font = BOLD
        s.cell(top + 1, 1, "WACC (rows) vs " + col_label + " (columns)")
        hdr_r = top + 2
        for j, f in enumerate(col_vals_formula):
            c = s.cell(hdr_r, 2 + j, f)
            c.number_format, c.font, c.fill = col_fmt, BOLD, HEAD_FILL
        for i, step in enumerate(config.SENS_WACC_STEPS):
            rr = hdr_r + 1 + i
            c = s.cell(rr, 1, f"={ref['wacc']}+({step})")
            c.number_format, c.font, c.fill = PCT2, BOLD, HEAD_FILL
            for j in range(len(col_vals_formula)):
                cl = get_column_letter(2 + j)
                cc = s.cell(rr, 2 + j, cell_formula(f"$A{rr}", f"{cl}${hdr_r}"))
                cc.number_format = USD
                if i == len(config.SENS_WACC_STEPS) // 2 and j == len(col_vals_formula) // 2:
                    cc.fill, cc.font = KEY_FILL, BOLD
        rng = f"B{hdr_r + 1}:{get_column_letter(1 + len(col_vals_formula))}{hdr_r + len(config.SENS_WACC_STEPS)}"
        s.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual",
                                                     formula=[ref["price"]],
                                                     font=Font(color="006100")))
        return hdr_r

    g_hdr = grid(3, "Gordon growth: implied value per share ($)", "terminal growth",
                 [f"={ref['tg']}+({st})" for st in config.SENS_GROWTH_STEPS],
                 lambda w, g: f"=IF({w}<={g},NA(),({pv_terms(w)}+'DCF'!${lc}${R_['fcf']}*(1+{g})/"
                              f"({w}-{g})/(1+{w})^{tvper}{bridge})/{ref['shares']})", PCT2)
    x_top = g_hdr + len(config.SENS_WACC_STEPS) + 3
    x_hdr = grid(x_top, "Exit multiple: implied value per share ($)", "exit EV/EBITDA",
                 [f"={ref['exit_multiple']}+({st})" for st in (-2, -1, 0, 1, 2)],
                 lambda w, x: f"=({pv_terms(w)}+'DCF'!${lc}${R_['ebitda']}*{x}/(1+{w})^{tvper}"
                              f"{bridge})/{ref['shares']}", MULT)
    s.cell(x_hdr + len(config.SENS_WACC_STEPS) + 2, 1,
           "Green text = at or above the current share price.")
    ref["sens_gordon_tl"] = (g_hdr + 1, 2)
    _widths(s, [16, 12, 12, 12, 12, 12])

    # ------------------------------------------------------------------ Drivers (Python output)
    dr = wb.create_sheet("Drivers")
    dr["A1"] = "What moves the answer most (one input moved at a time, full model re-run)"
    dr["A1"].font = TITLE
    dr["A2"] = ("Values computed by the Python model (dcf/core.py: tornado) and pasted as values; "
                "the base case ties to the DCF sheet (check below).")
    for j, hh in enumerate(["Driver", "Value at lower input ($)", "Value at higher input ($)",
                            "Swing ($)"], start=1):
        dr.cell(4, j, hh).font = BOLD
        dr.cell(4, j).fill = HEAD_FILL
    for i, t_ in enumerate(m.tornado):
        rr = 5 + i
        dr.cell(rr, 1, t_["driver"])
        for j, k in enumerate(["low_input_price", "high_input_price", "swing"], start=2):
            dr.cell(rr, j, round(t_[k], 4)).number_format = USD
    rr = 6 + len(m.tornado)
    dr.cell(rr, 1, "Base case (Python)")
    dr.cell(rr, 2, round(m.result.price_gordon, 6)).number_format = USD
    dr.cell(rr + 1, 1, "Base case (this workbook, DCF sheet)")
    dr.cell(rr + 1, 2, f"={ref['px_gordon']}").number_format = USD
    dr.cell(rr + 2, 1, "Difference")
    dr.cell(rr + 2, 2, f"=B{rr + 1}-B{rr}").number_format = '0.0000'
    br = rr + 4
    dr.cell(br, 1, "What would have to be true to justify the current share price").font = BOLD
    dr.cell(br, 1).fill = HEAD_FILL
    dr.cell(br + 1, 1, "(each input solved on its own, all others at base case)")
    labels_be = {"wacc": ("WACC", "wacc"), "ebit_margin": ("EBIT margin", "ebit_margin"),
                 "terminal_growth": ("Terminal growth", "terminal_growth"),
                 "growth_year1": ("Year-1 revenue growth (then fading)", "growth_year1")}
    for i, (k, (lab, fld)) in enumerate(labels_be.items()):
        dr.cell(br + 2 + i, 1, lab)
        v = m.breakeven.get(k)
        dr.cell(br + 2 + i, 2, round(v, 6) if v is not None else "not reachable").number_format = PCT2
        dr.cell(br + 2 + i, 3, f"base case: {getattr(m.assumptions, fld):.2%}")
    _widths(dr, [40, 26, 26, 12])

    # ------------------------------------------------------------------ Summary
    sm = summary
    sm["A1"] = f"{m.company} ({m.ticker}) - DCF valuation summary"
    sm["A1"].font = TITLE
    sm["A2"] = (f"Built from SEC EDGAR XBRL filings (CIK {m.cik}). Valuation date "
                f"{m.quote['date']}. $ millions except per-share data. Every cell below is a "
                "formula linked to the model sheets.")
    S = Rows(sm, start_row=4)
    S.header("Value per share")
    S.add("px", "Current share price ($)", f"={ref['price']}", USD)
    S.add("g", "DCF - Gordon growth ($)", f"={ref['px_gordon']}", USD, bold=True, fill=KEY_FILL)
    S.add("gu", "  upside / (downside) vs price", f"={S.at['g']}/{S.at['px']}-1", PCT)
    S.add("x", "DCF - exit multiple ($)", f"={ref['px_exit']}", USD, bold=True)
    S.add("xu", "  upside / (downside) vs price", f"={S.at['x']}/{S.at['px']}-1", PCT)
    S.add(None, "Trading comps - median peer EV/EBITDA ($)", f"={ref['comps_ev_med']}", USD)
    S.add(None, "Trading comps - median peer P/E ($)", f"={ref['comps_pe_med']}", USD)
    S.add(None, "52-week range of closing prices ($)",
          f'=TEXT({ref["week52_low"]},"0.00")&" - "&TEXT({ref["week52_high"]},"0.00")')
    S.gap()
    S.header("Key inputs")
    S.add(None, "WACC", f"={ref['wacc']}", PCT2)
    S.add(None, "Cost of equity", f"={ref['ke']}", PCT2)
    S.add(None, "Beta (5y monthly vs S&P 500)", f"={ref['beta']}", NUM3)
    S.add(None, "Risk-free rate", f"={ref['rf']}", PCT2)
    S.add(None, "Equity risk premium", f"={ref['erp']}", PCT2)
    S.add(None, "Terminal growth", f"={ref['tg']}", PCT2)
    S.add(None, "Exit EV/EBITDA (peer median)", f"={ref['exit_multiple']}", MULT)
    S.add(None, "Year-1 revenue growth", f"={ref['as_g1']}", PCT2)
    S.add(None, "EBIT margin", f"={ref['as_margin']}", PCT2)
    S.add(None, "Enterprise value (Gordon)", f"={ref['ev_gordon']}", MIL)
    S.add(None, "Terminal value share of EV (Gordon)", f"={ref['tvpct']}", PCT)
    S.gap()
    S.header("Sheets")
    for name, desc in [("Inputs", "market data, balance sheet, rates, rating table (sources cited)"),
                       ("Historicals", "6 fiscal years from 10-K XBRL with derived ratios"),
                       ("Beta", "60 monthly returns and the SLOPE regression"),
                       ("WACC", "CAPM cost of equity, synthetic-rating cost of debt, weights"),
                       ("Comps", "peer EV/EBITDA and P/E; sets the exit multiple"),
                       ("Assumptions", "every forecast driver, linked to its source"),
                       ("DCF", "5-year projection, discounting, both terminal values, equity bridge"),
                       ("Sensitivity", "WACC x terminal growth and WACC x exit multiple grids"),
                       ("Drivers", "one-at-a-time sensitivity (tornado) from the Python model")]:
        S.add(None, name, desc)
    _widths(sm, [44, 16, 70])

    wb.calculation.fullCalcOnLoad = True  # Excel recalculates every formula on open
    wb.save(path)
    return ref
