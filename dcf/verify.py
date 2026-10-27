"""Recalculate the generated workbook without Excel (pycel) so the formulas are checked, not
just written. openpyxl stores formulas only; pycel evaluates them the way Excel would."""
from __future__ import annotations

from openpyxl import load_workbook


def _find(ws, label: str, col: int = 2):
    for row in ws.iter_rows(min_col=1, max_col=1):
        if row[0].value == label:
            return f"{ws.title}!{chr(64 + col)}{row[0].row}"
    raise KeyError(label)


def _py314_shim() -> None:
    """pycel 1.0b30 still builds ast.Str nodes, which Python 3.14 removed. ast.Constant is the
    drop-in replacement (pycel only passes the string value)."""
    import ast
    if not hasattr(ast, "Str"):
        def Str(s=None, **kw):  # noqa: N802
            return ast.Constant(value=s if s is not None else kw.get("value"))
        ast.Str = Str


def recalc_workbook(path) -> dict:
    _py314_shim()
    from pycel import ExcelCompiler

    wb = load_workbook(path)
    cells = {
        "px_gordon": _find(wb["DCF"], "Implied value per share ($)", 2),
        "px_exit": _find(wb["DCF"], "Implied value per share ($)", 3),
        "wacc": _find(wb["WACC"], "WACC", 2),
        "beta": "Beta!H4",
    }
    xl = ExcelCompiler(filename=str(path), plugins=["dcf.pycel_extras"])
    return {k: float(xl.evaluate(v)) for k, v in cells.items()}


def evaluate(path, address: str):
    _py314_shim()
    from pycel import ExcelCompiler
    return ExcelCompiler(filename=str(path), plugins=["dcf.pycel_extras"]).evaluate(address)
