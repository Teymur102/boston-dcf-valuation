"""The Excel workbook must reach the same answer as the Python model.

Builds the model from the cached data in data/ (no network), writes the workbook to a temp
folder, recalculates every formula with pycel and compares with Python.
"""
import warnings

import numpy as np
import pytest
from openpyxl import load_workbook

from dcf import config, excel, pipeline
from dcf.verify import evaluate, recalc_workbook

pytestmark = pytest.mark.skipif(
    not (config.DATA_DIR / "sec" / "CIK0000109198.json.gz").exists(),
    reason="cached TJX data not present; run `python -m dcf run TJX` first")


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    warnings.filterwarnings("ignore")
    m = pipeline.build("TJX", config.DEFAULT_PEERS["TJX"])
    path = tmp_path_factory.mktemp("wb") / "TJX.xlsx"
    excel.write_workbook(m, path)
    return m, path


def test_workbook_price_matches_python(built):
    m, path = built
    x = recalc_workbook(path)
    assert x["px_gordon"] == pytest.approx(m.result.price_gordon, abs=1e-6)
    assert x["px_exit"] == pytest.approx(m.result.price_exit, abs=1e-6)
    assert x["wacc"] == pytest.approx(m.wacc, abs=1e-12)
    assert x["beta"] == pytest.approx(m.beta, abs=1e-12)


def test_every_sensitivity_cell_matches_python(built):
    m, path = built
    ws = load_workbook(path)["Sensitivity"]
    top = next(c.row for c in ws["A"] if c.value == "Gordon growth: implied value per share ($)")
    first = top + 3
    for i in range(len(m.sens_waccs)):
        for j in range(len(m.sens_growths)):
            addr = f"Sensitivity!{chr(66 + j)}{first + i}"
            assert evaluate(path, addr) == pytest.approx(m.sens_grid[i, j], abs=1e-6), addr


def test_core_cells_are_formulas_not_values(built):
    _, path = built
    wb = load_workbook(path)
    d = wb["DCF"]
    formulas = [c.value for row in d.iter_rows(min_row=4, max_row=24, min_col=3, max_col=7)
                for c in row if c.value is not None and not isinstance(c.value, int)]
    assert formulas and all(isinstance(v, str) and v.startswith("=") for v in formulas)


def test_comps_median_sets_exit_multiple(built):
    m, path = built
    ws = load_workbook(path)["Assumptions"]
    row = next(c.row for c in ws["A"] if c.value == "Exit EV/EBITDA multiple")
    v = evaluate(path, f"Assumptions!B{row}")
    assert v == pytest.approx(m.exit_multiple)
    assert v == pytest.approx(np.median([p.ev_ebitda for p in m.peers]))
